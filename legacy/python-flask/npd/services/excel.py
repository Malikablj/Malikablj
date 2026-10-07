"""Export ke Excel (.xlsx) memakai openpyxl."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="F2F2F7")


def build_workbook(sheets: list[dict]) -> bytes:
    """sheets = [{"name": "Projects", "rows": [[header...], [row...]], "widths": [14, 20, ...]}]"""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(title=sheet["name"][:31])
        for row in sheet["rows"]:
            ws.append(["" if v is None else v for v in row])
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        ws.freeze_panes = "A2"
        for i, width in enumerate(sheet.get("widths", []), start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
